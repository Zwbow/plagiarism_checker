"""
论文查重工具 - Web 版 (Streamlit App) - AI 语义查重终极版
运行方式：在终端执行 streamlit run app.py
"""

import os
# 必须在 import sentence_transformers 之前设置镜像环境变量
os.environ["HF_ENDPOINT"] = "https://hf-mirror.com"

import streamlit as st

# ================= 网页 UI 配置（必须是第一条 Streamlit 命令） =================
st.set_page_config(page_title="论文查重工具", page_icon="📄", layout="centered")

import jieba
import jieba.analyse
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import difflib
import re
import docx
import numpy as np
from sentence_transformers import SentenceTransformer

# ================= 核心算法与模型 =================

@st.cache_resource(show_spinner=False)
def load_model():
    """加载 AI 语义模型。优先本地，其次镜像站"""
    local_model_path = r"D:\new project\pythonProject\1\lunwenchachong\models"

    if os.path.exists(local_model_path):
        try:
            return SentenceTransformer(local_model_path)
        except Exception as e:
            st.warning(f"本地模型存在但加载失败: {e}，准备尝试联网下载...")

    st.warning("⚠️ 未找到本地模型或本地加载失败，正在回退到联网下载...")
    return SentenceTransformer('paraphrase-multilingual-MiniLM-L12-v2')

# 加载模型
model = None
try:
    with st.spinner("正在加载 AI 语义模型，如果首次运行需下载约 400MB 文件，请耐心等待..."):
        model = load_model()
except Exception as e:
    st.error("⚠️ AI 模型加载失败。这通常是因为网络问题，请检查网络后刷新重试。")
    st.exception(e)

def read_uploaded_file(uploaded_file):
    """从 Streamlit 上传的文件对象中读取文本，支持 .txt 和 .docx"""
    ext = os.path.splitext(uploaded_file.name)[1].lower()
    if ext == '.txt':
        return uploaded_file.read().decode('utf-8', errors='ignore')
    elif ext == '.docx':
        doc = docx.Document(uploaded_file)
        return '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
    else:
        st.warning(f"⚠️ 跳过不支持的文件格式：{uploaded_file.name}")
        return ''

def read_text(file_path):
    """通过路径读取本地文件"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    elif ext == '.docx':
        doc = docx.Document(file_path)
        return '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
    return ''

def collect_files(folder_path):
    """收集文件夹下所有文件的完整路径"""
    files = []
    for name in os.listdir(folder_path):
        full_path = os.path.join(folder_path, name)
        if os.path.isfile(full_path):
            files.append(full_path)
    return files

def cut_words(text):
    return ' '.join(jieba.lcut(text))

def calc_similarity(text_a, text_b):
    """传统 TF-IDF 字面相似度计算"""
    vectorizer = TfidfVectorizer()
    matrix = vectorizer.fit_transform([text_a, text_b])
    sim = cosine_similarity(matrix[0:1], matrix[1:2])
    return sim[0][0] * 100

def calc_semantic_similarity(text_a, text_b):
    """AI 语义相似度计算 (整篇文档级别，会被长文本稀释)"""
    if model is None or not text_a.strip() or not text_b.strip():
        return 0.0
    embeddings = model.encode([text_a, text_b], convert_to_numpy=True)
    norm_a = np.linalg.norm(embeddings[0])
    norm_b = np.linalg.norm(embeddings[1])
    if norm_a == 0 or norm_b == 0:
        return 0.0
    cosine = np.dot(embeddings[0], embeddings[1]) / (norm_a * norm_b)
    return float(cosine) * 100

def split_sentences(text):
    """把长文本按句号、问号、感叹号、换行切分成句子列表"""
    parts = re.split(r'(?<=[。！？\n])', text)
    return [s.strip() for s in parts if len(s.strip()) > 5]

def find_duplicate_sentences_from_lists(sentences_a, sentences_b, threshold=0.8):
    """【字面级别】逐句比对（基于 difflib 字面相似度）"""
    duplicates = []
    for s_a in sentences_a:
        for s_b in sentences_b:
            ratio = difflib.SequenceMatcher(None, s_a, s_b).ratio()
            if ratio >= threshold:
                duplicates.append((s_a, s_b, ratio))
                break
    return duplicates

def find_duplicate_sentences_ai(sentences_a, sentences_b, threshold=0.5):
    """【AI语义级别】逐句比对（基于 Sentence-BERT 向量相似度）"""
    if not sentences_a or not sentences_b or model is None:
        return []

    # 批量把句子列表转成向量矩阵（比单句循环快很多）
    embeddings_a = model.encode(sentences_a, convert_to_numpy=True)
    embeddings_b = model.encode(sentences_b, convert_to_numpy=True)

    duplicates = []
    for i, emb_a in enumerate(embeddings_a):
        norm_a = np.linalg.norm(emb_a)
        if norm_a == 0:
            continue
        best_score = 0.0
        best_match = None
        for j, emb_b in enumerate(embeddings_b):
            norm_b = np.linalg.norm(emb_b)
            if norm_b == 0:
                continue
            cosine = np.dot(emb_a, emb_b) / (norm_a * norm_b)
            score = float(cosine)
            if score > best_score:
                best_score = score
                best_match = sentences_b[j]
        if best_score >= threshold:
            duplicates.append((sentences_a[i], best_match, best_score))
    return duplicates

# ================= 网页 UI 与 主逻辑 =================

st.title("📄 论文查重工具 (AI 语义版)")
st.markdown("上传你的论文和比对文献库，一键生成查重报告。支持传统字面查重与 AI 语义查重。")

if model is not None:
    st.subheader("1. 上传待查论文")
    target_file = st.file_uploader("请上传你要查重的论文（支持 .txt, .docx）", type=['txt', 'docx'], key="target")

    st.subheader("2. 选择比对文献库")
    input_mode = st.radio("请选择比对文献库的输入方式：", ["网页上传文件", "输入本地文件夹路径"])

    source_texts = {}

    if input_mode == "网页上传文件":
        source_files = st.file_uploader(
            "请上传要比对的文献库（可多选，支持 .txt, .docx）",
            type=['txt', 'docx'],
            accept_multiple_files=True,
            key="sources"
        )
        if source_files:
            for sf in source_files:
                source_texts[sf.name] = read_uploaded_file(sf)

    elif input_mode == "输入本地文件夹路径":
        folder_path = st.text_input("请输入本地文件夹的绝对路径（例如 D:\\new project\\...\\datas）：")
        if folder_path:
            if os.path.isdir(folder_path):
                local_files = collect_files(folder_path)
                for path in local_files:
                    source_texts[os.path.basename(path)] = read_text(path)
                st.success(f"✅ 成功读取文件夹，共找到 {len(source_texts)} 个文件")
            else:
                st.error("❌ 文件夹路径无效，请检查！")

    # 3. 开始查重
    if st.button("🚀 开始查重", use_container_width=True):
        if not target_file:
            st.error("❌ 请先上传待查论文！")
        elif not source_texts:
            st.error("❌ 请至少提供一篇比对文献（上传文件或输入文件夹路径）！")
        else:
            with st.spinner("正在比对中，请稍候..."):
                target_text = read_uploaded_file(target_file)
                target_cut = cut_words(target_text)

                # 预计算目标文本的关键词和句子（只算一次）
                target_keywords = set(jieba.analyse.extract_tags(target_text, topK=8))
                target_sentences = split_sentences(target_text)

                results = {}
                report_data = {}
                progress_bar = st.progress(0)
                items = list(source_texts.items())

                for idx, (file_name, other_text) in enumerate(items):
                    if not other_text.strip():
                        continue

                    other_cut = cut_words(other_text)
                    other_sentences = split_sentences(other_text)

                    # === 字面相似度（整篇 TF-IDF） ===
                    score_tfidf = calc_similarity(target_cut, other_cut)

                    # === AI 逐句比对（最重要的一步）===
                    dup_ai = find_duplicate_sentences_ai(target_sentences, other_sentences, threshold=0.5)

                    # === 字面逐句比对（作为补充）===
                    dup_tfidf = find_duplicate_sentences_from_lists(target_sentences, other_sentences, threshold=0.8)

                    # === 合并去重，按相似度从高到低排序 ===
                    seen = set()
                    merged = []
                    for s_a, s_b, ratio in dup_ai + dup_tfidf:
                        key = (s_a, s_b)
                        if key not in seen:
                            seen.add(key)
                            merged.append((s_a, s_b, ratio))
                    merged.sort(key=lambda x: x[2], reverse=True)
                    dup_sentences = merged

                    # === 综合 AI 得分：如果存在逐句匹配，取最高分；否则用整篇文档相似度 ===
                    if dup_ai:
                        score_ai = max(d[2] for d in dup_ai) * 100
                    else:
                        score_ai = calc_semantic_similarity(target_text, other_text)

                    # 最终综合得分 = 两者取高
                    score = max(score_tfidf, score_ai)

                    results[file_name] = score

                    # 关键词交集
                    other_keywords = set(jieba.analyse.extract_tags(other_text, topK=8))
                    common = list(target_keywords & other_keywords)

                    report_data[file_name] = {
                        'score': score,
                        'score_tfidf': score_tfidf,
                        'score_ai': score_ai,
                        'common_words': common,
                        'duplicates': dup_sentences
                    }

                    progress_bar.progress((idx + 1) / len(items))

                # ===== 展示网页端结果 =====
                st.success("✅ 查重完成！")

                if results:
                    ranked = sorted(results.items(), key=lambda x: x[1], reverse=True)

                    st.subheader("📊 查重率极值报告")
                    col1, col2 = st.columns(2)
                    col1.metric("🔴 最高相似度", f"{ranked[0][1]:.2f}%", ranked[0][0])
                    col2.metric("🟢 最低相似度", f"{ranked[-1][1]:.2f}%", ranked[-1][0])

                    st.markdown("### 📄 各文件详细结果")
                    for name, data in sorted(report_data.items(), key=lambda x: x[1]['score'], reverse=True):
                        score = data['score']
                        with st.expander(f"📄 {name} —— 综合相似度：{score:.2f}%"):
                            st.markdown(f"**🤖 AI语义相似度：** {data['score_ai']:.2f}% （传统字面相似度：{data['score_tfidf']:.2f}%）")
                            st.markdown(
                                f"**📌 高频重复词汇：** {', '.join(data['common_words']) if data['common_words'] else '无'}")

                            if data['duplicates']:
                                st.markdown("**🔍 疑似抄袭句子（AI 语义 + 字面合并）：**")
                                for s_a, s_b, ratio in data['duplicates']:
                                    st.warning(f"**原文：** {s_a}\n\n**对比：** {s_b}\n\n*相似度：{ratio * 100:.1f}%*")
                            else:
                                st.success("✅ 未发现高度重复的句子。")

                    # ===== 生成 HTML 报告用于下载 =====
                    html = '''<!DOCTYPE html><html><head><meta charset="utf-8"><title>查重报告</title>
                    <style>body{font-family:"微软雅黑",sans-serif;padding:30px;background:#f8f9fa;color:#333;}
                    .container{max-width:900px;margin:0 auto;background:white;padding:30px;border-radius:10px;box-shadow:0 4px 10px rgba(0,0,0,0.1);}
                    h1{text-align:center;color:#2c3e50;}.summary{background:#e8f4fd;padding:20px;border-radius:8px;margin-bottom:30px;}
                    .file-card{border:1px solid #e0e0e0;border-radius:8px;margin-bottom:20px;overflow:hidden;}
                    .file-header{background:#f1f3f5;padding:15px;font-weight:bold;display:flex;justify-content:space-between;}
                    .file-body{padding:15px;}.high{color:#e74c3c;font-weight:bold;}.mid{color:#e67e22;}.low{color:#27ae60;}
                    .dup-box{background:#fff3cd;border-left:4px solid #ffc107;padding:10px;margin:10px 0;border-radius:4px;}
                    .dup-box p{margin:5px 0;}.match{background-color:#ffcccc;padding:2px 4px;border-radius:3px;}</style>
                    </head><body><div class="container"><h1>📄 论文查重报告 (AI 语义版)</h1><div class="summary">
                    <p><b>待查文件：</b>''' + target_file.name + '''</p>
                    <p><b>比对文件总数：</b>''' + str(len(report_data)) + ''' 个</p>
                    <p><b>🔴 最高相似度：</b>''' + ranked[0][0] + ' —— ' + f'{ranked[0][1]:.2f}%' + '''</p>
                    <p><b>🟢 最低相似度：</b>''' + ranked[-1][0] + ' —— ' + f'{ranked[-1][1]:.2f}%' + '''</p></div>'''

                    for name, data in sorted(report_data.items(), key=lambda x: x[1]['score'], reverse=True):
                        score_color = 'high' if data['score'] > 50 else 'mid' if data['score'] > 20 else 'low'
                        html += f'''<div class="file-card"><div class="file-header"><span>📄 {name}</span>
                        <span class="{score_color}">综合相似度：{data['score']:.2f}%</span></div><div class="file-body">'''
                        html += f"<p><b>🤖 AI语义相似度：</b>{data['score_ai']:.2f}% （传统字面相似度：{data['score_tfidf']:.2f}%）</p>"
                        if data['common_words']:
                            html += f"<p><b>📌 高频重复词汇：</b>{', '.join(data['common_words'])}</p>"
                        if data['duplicates']:
                            html += "<p><b>🔍 疑似抄袭句子：</b></p>"
                            for s_a, s_b, ratio in data['duplicates']:
                                html += f'''<div class="dup-box"><p>📝 原文：<span class="match">{s_a}</span></p>
                                <p>📄 对比：<span class="match">{s_b}</span></p>
                                <p style="color:#888;font-size:12px;">相似度：{ratio * 100:.1f}%</p></div>'''
                        else:
                            html += "<p>✅ 未发现高度重复的句子。</p>"
                        html += "</div></div>"
                    html += "</div></body></html>"

                    st.download_button(
                        label="📥 下载 HTML 详细报告",
                        data=html,
                        file_name="查重报告.html",
                        mime="text/html",
                        use_container_width=True
                    )
                else:
                    st.warning("❌ 没有可比较的文件。")